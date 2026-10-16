"""Human review and approved document lifecycle."""
import json
from pathlib import Path
from docx import Document
from ..repositories import sqlite as repository
from ..core.settings import settings
from .processing import set_status


class ReviewError(Exception):
    def __init__(self, status_code: int, message: str):
        super().__init__(message)
        self.status_code = status_code


def pending_draft(draft_id: str) -> dict:
    draft = repository.one("SELECT * FROM drafts WHERE id=?", (draft_id,))
    if not draft:
        raise ReviewError(404, "Черновик не найден")
    if draft["status"] != "PENDING":
        raise ReviewError(409, "Черновик уже рассмотрен")
    return draft


def edit_draft(draft_id: str, text: str, actor: str) -> dict:
    draft = pending_draft(draft_id)
    repository.execute("UPDATE drafts SET edited_text=? WHERE id=?", (text, draft_id))
    repository.audit(draft["appeal_id"], "draft_edited", actor, {"draft_id": draft_id})
    return {"status": "saved"}


def review_draft(draft_id: str, actor: str, text: str | None, state: str) -> dict:
    draft = pending_draft(draft_id)
    final = text or draft["edited_text"]
    repository.execute("UPDATE drafts SET edited_text=?,status=?,reviewer=?,reviewed_at=? WHERE id=?",
                       (final, state, actor, repository.now(), draft_id))
    set_status(draft["appeal_id"], "APPROVED" if state == "APPROVED" else "REJECTED")
    repository.audit(draft["appeal_id"], "draft_" + state.lower(), actor,
                     {"draft_id": draft_id, "changed": final != json.loads(draft["ai_payload"])["full_text"]})
    return {"status": state, "draft_id": draft_id}


def export_docx(draft_id: str) -> Path:
    draft = repository.one("SELECT * FROM drafts WHERE id=?", (draft_id,))
    if not draft:
        raise ReviewError(404, "Черновик не найден")
    if draft["status"] != "APPROVED":
        raise ReviewError(409, "Экспорт итогового ответа возможен после одобрения")
    path = Path(settings.storage_base_path) / f"{draft_id}.docx"
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading("Ответ на обращение", level=1)
    for paragraph in draft["edited_text"].split("\n"):
        document.add_paragraph(paragraph)
    document.save(path)
    repository.execute("UPDATE drafts SET document_uri=? WHERE id=?", (str(path), draft_id))
    repository.audit(draft["appeal_id"], "docx_exported", draft["reviewer"], {"draft_id": draft_id})
    return path
