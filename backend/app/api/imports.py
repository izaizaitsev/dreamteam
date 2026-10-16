"""Imports API routes."""
import json
import os
from io import BytesIO
from pathlib import Path
from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse
from docx import Document
from ..repositories import sqlite as s
from ..services.importers import BitrixJsonAdapter, records, normalize_text
from ..services.indexing import index_case
from ..core.settings import settings
from .dependencies import auth, admin, fail

router = APIRouter()


@router.post("/api/import/appeals", dependencies=[Depends(admin)])
async def import_appeals(request: Request):
    try:
        items = records(await request.json())
    except (ValueError, json.JSONDecodeError) as exc:
        fail(exc)
    report = {"total": len(items), "created": 0, "updated": 0, "skipped": 0, "errors": []}
    adapter = BitrixJsonAdapter()
    for index, item in enumerate(items):
        try:
            dto = adapter.convert(item)
            existing = s.one("SELECT id,text,bitrix_status FROM appeals WHERE external_id=?", (dto.external_id,))
            if existing and existing["text"] == dto.text and existing["bitrix_status"] == dto.status:
                report["skipped"] += 1
                continue
            if existing:
                s.execute("UPDATE appeals SET source=?,title=?,text=?,applicant_name=?,applicant_type=?,building_id=?,apartment_id=?,received_at=?,due_at=?,bitrix_status=?,raw_payload=?,updated_at=? WHERE id=?",
                    (dto.source, dto.title, dto.text, dto.applicant_name, dto.applicant_type,
                     dto.building_id, dto.apartment_id, str(dto.received_at) if dto.received_at else None,
                     str(dto.due_at) if dto.due_at else None, dto.status,
                     json.dumps(dto.raw_payload, ensure_ascii=False), s.now(), existing["id"]))
                report["updated"] += 1
                s.audit(existing["id"], "appeal_updated")
            else:
                new_id = s.uid()
                s.execute("INSERT INTO appeals VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (new_id, dto.external_id, dto.source, dto.title, dto.text, dto.applicant_name,
                     dto.applicant_type, dto.building_id, dto.apartment_id,
                     str(dto.received_at) if dto.received_at else None,
                     str(dto.due_at) if dto.due_at else None, dto.status,
                     json.dumps(dto.raw_payload, ensure_ascii=False), "NEW", None, s.now(), s.now()))
                report["created"] += 1
                s.audit(new_id, "appeal_imported")
        except Exception as exc:
            report["errors"].append({"row": index + 1, "message": str(exc)})
    return report



@router.post("/api/import/resolved-cases", dependencies=[Depends(admin)])
async def import_cases(request: Request, index_now: bool = False):
    try:
        items = records(await request.json())
    except (ValueError, json.JSONDecodeError) as exc:
        fail(exc)
    report = {"total": len(items), "created": 0, "updated": 0, "errors": [], "indexed": 0}
    for index, item in enumerate(items):
        try:
            external = str(item.get("external_id") or item.get("ID") or "").strip()
            text = normalize_text(item.get("appeal_text") or item.get("text") or "")
            if not external or not text:
                raise ValueError("Нужны external_id и appeal_text")
            existing = s.one("SELECT id,decision_text,decision_file_uri FROM cases WHERE external_id=?", (external,))
            case_id = existing["id"] if existing else s.uid()
            decision_text = item.get("decision_text") if "decision_text" in item else (existing["decision_text"] if existing else None)
            decision_file_uri = item.get("decision_file_uri") or (existing["decision_file_uri"] if existing else None)
            stored_path = Path(settings.storage_base_path) / "resolved" / f"{case_id}.docx"
            if existing and stored_path.is_file():
                # The uploaded Word document remains the source of the filled response.
                decision_text = existing["decision_text"]
                decision_file_uri = str(stored_path)
            values = (text, decision_text, decision_file_uri,
                      item.get("category_code"), item.get("source"), item.get("building_id"),
                      json.dumps(item, ensure_ascii=False), s.now())
            if existing:
                s.execute("UPDATE cases SET appeal_text=?,decision_text=?,decision_file_uri=?,category_code=?,source=?,building_id=?,metadata=?,updated_at=?,indexed_at=NULL WHERE id=?",
                          (*values, case_id))
                report["updated"] += 1
            else:
                s.execute("INSERT INTO cases VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (case_id, external, text, item.get("decision_text"), item.get("decision_file_uri"),
                     item.get("category_code"), item.get("source"), item.get("building_id"),
                     json.dumps(item, ensure_ascii=False), None, None, s.now(), s.now()))
                report["created"] += 1
            if index_now:
                try:
                    await index_case(case_id)
                    report["indexed"] += 1
                except Exception as exc:
                    report["errors"].append({"row": index + 1, "message": "Индексация: " + str(exc)})
        except Exception as exc:
            report["errors"].append({"row": index + 1, "message": str(exc)})
    return report



@router.post("/api/cases/decision-file", dependencies=[Depends(admin)])
async def upload_decision_file(external_id: str, file: UploadFile = File(...)):
    case = s.one("SELECT id FROM cases WHERE external_id=?", (external_id,))
    if not case:
        raise HTTPException(404, "Сначала импортируйте решённое обращение с этим external_id")
    if not file.filename or not file.filename.lower().endswith(".docx"):
        raise HTTPException(400, "Нужен файл .docx")
    content = await file.read(10 * 1024 * 1024 + 1)
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(413, "Файл больше 10 МБ")
    try:
        document = Document(BytesIO(content))
        parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text.strip() for cell in row.cells if cell.text.strip())
        decision_text = "\n".join(parts)
    except Exception as exc:
        raise HTTPException(400, "Не удалось прочитать Word-файл") from exc
    if not decision_text:
        raise HTTPException(400, "В Word-файле нет текстового решения")
    path = Path(settings.storage_base_path) / "resolved" / f"{case['id']}.docx"
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp")
    try:
        temporary.write_bytes(content)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)
    s.execute("UPDATE cases SET decision_text=?,decision_file_uri=?,updated_at=? WHERE id=?",
              (decision_text, str(path), s.now(), case["id"]))
    s.audit(None, "decision_file_uploaded", details={"case_id": case["id"]})
    return {"case_id": case["id"], "external_id": external_id, "decision_file_uri": str(path),
            "extracted_characters": len(decision_text)}



@router.get("/api/cases/{case_id}/decision-file", dependencies=[Depends(auth)])
def download_decision_file(case_id: str):
    if not s.one("SELECT id FROM cases WHERE id=?", (case_id,)):
        raise HTTPException(404, "Кейс не найден")
    path = Path(settings.storage_base_path) / "resolved" / f"{case_id}.docx"
    if not path.is_file():
        raise HTTPException(404, "Word-файл для этого кейса не загружен")
    return FileResponse(path, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        filename=f"resolved-{case_id}.docx")
